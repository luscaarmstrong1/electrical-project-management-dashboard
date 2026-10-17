from __future__ import annotations

import random
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd


SEED = 20260908
REFERENCE_DATE = date(2026, 9, 8)
ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"

LOCALIZE_VALUES = {
    "Industria Alfa": "Indústria Alfa",
    "Centro Logistico Delta": "Centro Logístico Delta",
    "Condominio Solar Norte": "Condomínio Solar Norte",
    "Saude": "Saúde",
    "Metalurgica Orion": "Metalúrgica Orion",
    "Agropecuaria Vale Verde": "Agropecuária Vale Verde",
    "Shopping Estacao": "Shopping Estação",
    "Educacao": "Educação",
    "Logistica": "Logística",
    "Mineracao": "Mineração",
    "Textil Primavera": "Têxtil Primavera",
    "Laticinios Boa Fonte": "Laticínios Boa Fonte",
    "Terminal Rodoviario Central": "Terminal Rodoviário Central",
    "Ceramica Horizonte": "Cerâmica Horizonte",
    "Servicos": "Serviços",
    "Gestor de Operacoes": "Gestor de Operações",
    "Assistente Tecnico": "Assistente Técnico",
    "Operacoes": "Operações",
    "EDP Sao Paulo": "EDP São Paulo",
    "EDP Espirito Santo": "EDP Espírito Santo",
    "Enel Sao Paulo": "Enel São Paulo",
    "Enel Ceara": "Enel Ceará",
    "Geracao distribuida": "Geração distribuída",
    "Microgeracao Distribuida": "Microgeração Distribuída",
    "Minigeracao Distribuida": "Minigeração Distribuída",
    "Ampliacao de Geracao": "Ampliação de Geração",
    "Infraestrutura eletrica": "Infraestrutura elétrica",
    "Subestacao": "Subestação",
    "Adequacao Eletrica": "Adequação Elétrica",
    "Projeto de Protecao": "Projeto de Proteção",
    "Protecao e estudos": "Proteção e estudos",
    "Estudo de Coordenacao e Seletividade": "Estudo de Coordenação e Seletividade",
    "Mobilidade eletrica": "Mobilidade elétrica",
    "Eficiencia Energetica": "Eficiência Energética",
    "Energia e operacoes": "Energia e operações",
    "Automacao Industrial": "Automação Industrial",
    "Automacao": "Automação",
    "Migracao Mercado Livre": "Migração Mercado Livre",
    "Gestao de Energia": "Gestão de Energia",
    "Infraestrutura Eletrica Industrial": "Infraestrutura Elétrica Industrial",
    "Nao iniciado": "Não iniciado",
    "Em analise": "Em análise",
    "Em aprovacao": "Em aprovação",
    "Aguardando concessionaria": "Aguardando concessionária",
    "Em execucao": "Em execução",
    "Concluido": "Concluído",
    "CONCLUIDO": "CONCLUÍDO",
    "Analise tecnica": "Análise técnica",
    "Analise": "Análise",
    "Aprovacao": "Aprovação",
    "Adequacoes": "Adequações",
    "Execucao": "Execução",
    "Conclusao": "Conclusão",
    "Media": "Média",
    "Critica": "Crítica",
    "Encerrado sem acao": "Encerrado sem ação",
    "Documento cliente": "Documento do cliente",
    "Informacao tecnica": "Informação técnica",
    "Aprovacao interna": "Aprovação interna",
    "Concessionaria": "Concessionária",
    "Alteracao de projeto": "Alteração de projeto",
    "Pendencia sintetica para demonstracao de gestao operacional": "Pendência sintética para demonstração de gestão operacional",
    "Revisao": "Revisão",
    "Gestao": "Gestão",
    "Dependencia externa da distribuidora": "Dependência externa da distribuidora",
    "Aguardando documentos ou decisao do cliente": "Aguardando documentos ou decisão do cliente",
    "Projeto com pendencias registradas": "Projeto com pendências registradas",
    "Projeto encerrado": "Projeto encerrado",
    "Concluido no prazo": "Concluído no prazo",
    "Fluxo regular": "Fluxo regular",
    "Mais de 30 dias": "Mais de 30 dias",
}


def brl_round(value: float) -> float:
    return round(float(value), 2)


def weighted_choice(items: list[str], weights: list[float]) -> str:
    return random.choices(items, weights=weights, k=1)[0]


def build_dimensions() -> dict[str, pd.DataFrame]:
    clients = [
        "Industria Alfa", "Comercial Beta", "Grupo Horizonte", "Fazenda Aurora",
        "Centro Logistico Delta", "Condominio Solar Norte", "Hospital Vera Cruz",
        "Rede Varejo Prisma", "Metalurgica Orion", "Agropecuaria Vale Verde",
        "Shopping Estacao", "Cooperativa Campo Sul", "Frigorifico Atlas",
        "Universidade Nova Energia", "Porto Seco Leste", "Mineracao Serra Azul",
        "Textil Primavera", "Laticinios Boa Fonte", "Data Center Sigma",
        "Usina BioSol", "Terminal Rodoviario Central", "Ceramica Horizonte",
        "Laboratorio Innova", "Hotel Mirante", "Aeroporto Regional"
    ]
    dim_cliente = pd.DataFrame({
        "ID_Cliente": [f"CLI{idx:03d}" for idx in range(1, len(clients) + 1)],
        "Cliente": clients,
        "Segmento": [
            "Industrial", "Comercial", "Holding", "Agro", "Logistica",
            "Residencial coletivo", "Saude", "Varejo", "Industrial", "Agro",
            "Comercial", "Agro", "Industrial", "Educacao", "Logistica",
            "Mineracao", "Industrial", "Alimentos", "Tecnologia", "Energia",
            "Transporte", "Industrial", "Servicos", "Hotelaria", "Transporte"
        ],
    })

    responsaveis = [
        ("Lucas Silva Santos", "Engenheiro de Projetos"),
        ("Mariana Costa", "Engenheira Eletricista"),
        ("Rafael Martins", "Analista de Projetos"),
        ("Gabriel Oliveira", "Projetista"),
        ("Camila Souza", "Coordenadora"),
        ("Fernando Lima", "Gestor de Operacoes"),
        ("Bruno Rocha", "Assistente Tecnico"),
        ("Ana Ribeiro", "Engenheira Eletricista"),
        ("Patricia Nunes", "Projetista"),
        ("Diego Carvalho", "Analista de Projetos"),
    ]
    dim_responsavel = pd.DataFrame({
        "ID_Responsavel": [f"RESP{idx:02d}" for idx in range(1, len(responsaveis) + 1)],
        "Responsavel": [r[0] for r in responsaveis],
        "Cargo": [r[1] for r in responsaveis],
        "Squad": ["Projetos", "Engenharia", "PMO", "Desenho", "Operacoes", "Operacoes", "Campo", "Engenharia", "Desenho", "PMO"],
    })

    concessionarias = [
        "CEMIG", "CPFL Paulista", "CPFL Piratininga", "Neoenergia Elektro",
        "EDP Sao Paulo", "EDP Espirito Santo", "Enel Sao Paulo", "Enel Ceara",
        "Energisa", "Equatorial", "Celesc", "Outras"
    ]
    dim_concessionaria = pd.DataFrame({
        "ID_Concessionaria": [f"CONC{idx:02d}" for idx in range(1, len(concessionarias) + 1)],
        "Concessionaria": concessionarias,
    })

    tipos = [
        ("Sistema Fotovoltaico", "Geracao distribuida", "BT"),
        ("Microgeracao Distribuida", "Geracao distribuida", "BT"),
        ("Minigeracao Distribuida", "Geracao distribuida", "MT"),
        ("Ampliacao de Geracao", "Geracao distribuida", "MT"),
        ("Entrada de Energia BT", "Entrada de energia", "BT"),
        ("Entrada de Energia MT", "Entrada de energia", "MT"),
        ("Subestacao", "Infraestrutura eletrica", "MT"),
        ("Adequacao Eletrica", "Infraestrutura eletrica", "BT"),
        ("Projeto de Protecao", "Protecao e estudos", "MT"),
        ("Estudo de Coordenacao e Seletividade", "Protecao e estudos", "MT"),
        ("Eletroposto", "Mobilidade eletrica", "BT"),
        ("Eficiencia Energetica", "Energia e operacoes", "BT"),
        ("Automacao Industrial", "Automacao", "BT"),
        ("Migracao Mercado Livre", "Energia e operacoes", "MT"),
        ("Gestao de Energia", "Energia e operacoes", "BT"),
        ("Infraestrutura Eletrica Industrial", "Infraestrutura eletrica", "MT"),
    ]
    dim_tipo = pd.DataFrame({
        "ID_TipoProjeto": [f"TIPO{idx:02d}" for idx in range(1, len(tipos) + 1)],
        "Tipo_Projeto": [t[0] for t in tipos],
        "Categoria_Engenharia": [t[1] for t in tipos],
        "Nivel_Tensao_Padrao": [t[2] for t in tipos],
    })

    statuses = [
        "Nao iniciado", "Em analise", "Em desenvolvimento", "Em aprovacao",
        "Aguardando cliente", "Aguardando concessionaria", "Em execucao",
        "Em comissionamento", "Concluido", "Suspenso", "Cancelado"
    ]
    dim_status = pd.DataFrame({
        "ID_Status": [f"STAT{idx:02d}" for idx in range(1, len(statuses) + 1)],
        "Status": statuses,
        "Status_Simplificado_Base": [
            "ABERTO", "ABERTO", "ABERTO", "ABERTO", "ABERTO", "ABERTO",
            "ABERTO", "ABERTO", "CONCLUIDO", "SUSPENSO", "CANCELADO"
        ],
    })

    etapas = [
        "Levantamento inicial", "Analise tecnica", "Coleta de documentos",
        "Desenvolvimento do projeto", "Revisao interna", "Protocolo",
        "Analise da concessionaria", "Adequacoes", "Aprovacao", "Execucao",
        "Vistoria", "Comissionamento", "Encerramento"
    ]
    dim_etapa = pd.DataFrame({
        "ID_Etapa": [f"ETAPA{idx:02d}" for idx in range(1, len(etapas) + 1)],
        "Ordem_Etapa": list(range(1, len(etapas) + 1)),
        "Etapa": etapas,
        "Macro_Etapa": [
            "Entrada", "Analise", "Analise", "Desenvolvimento", "Desenvolvimento",
            "Protocolo", "Aprovacao", "Aprovacao", "Aprovacao", "Execucao",
            "Execucao", "Conclusao", "Conclusao"
        ],
    })

    localidades = [
        ("SP", "São Paulo", "Sudeste", 0.34), ("MG", "Belo Horizonte", "Sudeste", 0.22),
        ("ES", "Vitória", "Sudeste", 0.12), ("GO", "Goiânia", "Centro-Oeste", 0.07),
        ("MT", "Cuiabá", "Centro-Oeste", 0.05), ("MS", "Campo Grande", "Centro-Oeste", 0.04),
        ("PR", "Curitiba", "Sul", 0.05), ("SC", "Joinville", "Sul", 0.04),
        ("RJ", "Rio de Janeiro", "Sudeste", 0.04), ("CE", "Fortaleza", "Nordeste", 0.03),
    ]
    dim_localidade = pd.DataFrame({
        "ID_Localidade": [f"LOC{idx:02d}" for idx in range(1, len(localidades) + 1)],
        "Estado": [l[0] for l in localidades],
        "Municipio": [l[1] for l in localidades],
        "Regiao": [l[2] for l in localidades],
        "Peso_Carteira": [l[3] for l in localidades],
    })

    calendario = pd.DataFrame({"Data": pd.date_range("2024-01-01", "2027-12-31", freq="D")})
    calendario["Ano"] = calendario["Data"].dt.year
    calendario["Mes"] = calendario["Data"].dt.month_name(locale=None).str[:3]
    calendario["Numero_Mes"] = calendario["Data"].dt.month
    calendario["Mes_Ano"] = calendario["Data"].dt.strftime("%b/%Y")
    calendario["Trimestre"] = "T" + calendario["Data"].dt.quarter.astype(str)
    calendario["Semana"] = calendario["Data"].dt.isocalendar().week.astype(int)
    calendario["Dia"] = calendario["Data"].dt.day
    calendario["Dia_Semana"] = calendario["Data"].dt.day_name(locale=None)
    calendario["Ano_Mes"] = calendario["Data"].dt.strftime("%Y-%m")
    calendario["Inicio_Mes"] = calendario["Data"].values.astype("datetime64[M]")
    calendario["Fim_Mes"] = calendario["Data"] + pd.offsets.MonthEnd(0)
    calendario["Data"] = calendario["Data"].dt.date
    calendario["Inicio_Mes"] = pd.to_datetime(calendario["Inicio_Mes"]).dt.date
    calendario["Fim_Mes"] = pd.to_datetime(calendario["Fim_Mes"]).dt.date

    return {
        "Dim_Cliente": dim_cliente,
        "Dim_Responsavel": dim_responsavel,
        "Dim_Concessionaria": dim_concessionaria,
        "Dim_TipoProjeto": dim_tipo,
        "Dim_Status": dim_status,
        "Dim_Etapa": dim_etapa,
        "Dim_Localidade": dim_localidade,
        "Dim_Calendario": calendario,
    }


def project_duration_days(tipo: str, complexity: str) -> int:
    base = {
        "Sistema Fotovoltaico": 75, "Microgeracao Distribuida": 65,
        "Minigeracao Distribuida": 115, "Ampliacao de Geracao": 125,
        "Entrada de Energia BT": 45, "Entrada de Energia MT": 80,
        "Subestacao": 150, "Adequacao Eletrica": 55,
        "Projeto de Protecao": 70, "Estudo de Coordenacao e Seletividade": 60,
        "Eletroposto": 110, "Eficiencia Energetica": 50,
        "Automacao Industrial": 95, "Migracao Mercado Livre": 85,
        "Gestao de Energia": 45, "Infraestrutura Eletrica Industrial": 130,
    }[tipo]
    factor = {"Baixa": 0.85, "Media": 1.0, "Alta": 1.25, "Critica": 1.45}[complexity]
    return int(base * factor + random.randint(-12, 18))


def power_for_type(tipo: str) -> float:
    ranges = {
        "Sistema Fotovoltaico": (8, 250), "Microgeracao Distribuida": (5, 75),
        "Minigeracao Distribuida": (80, 1500), "Ampliacao de Geracao": (50, 900),
        "Entrada de Energia BT": (5, 60), "Entrada de Energia MT": (75, 600),
        "Subestacao": (300, 3500), "Adequacao Eletrica": (10, 180),
        "Projeto de Protecao": (75, 1200), "Estudo de Coordenacao e Seletividade": (75, 1600),
        "Eletroposto": (60, 450), "Eficiencia Energetica": (10, 300),
        "Automacao Industrial": (20, 700), "Migracao Mercado Livre": (300, 3000),
        "Gestao de Energia": (20, 500), "Infraestrutura Eletrica Industrial": (150, 2500),
    }[tipo]
    return round(float(np.random.triangular(ranges[0], (ranges[0] + ranges[1]) / 3, ranges[1])), 1)


def create_projects(dims: dict[str, pd.DataFrame], n: int = 200) -> pd.DataFrame:
    tipos_weights = [13, 11, 9, 5, 8, 7, 6, 8, 5, 4, 5, 5, 5, 4, 3, 7]
    states = dims["Dim_Localidade"]["Estado"].tolist()
    state_weights = dims["Dim_Localidade"]["Peso_Carteira"].tolist()
    conc_weights = [17, 14, 8, 10, 11, 7, 12, 3, 6, 4, 5, 3]
    status_weights = [5, 9, 17, 10, 9, 12, 12, 5, 16, 3, 2]
    priorities = ["Baixa", "Media", "Alta", "Critica"]
    complexities = ["Baixa", "Media", "Alta", "Critica"]
    project_prefixes = ["Solaris", "Atlas", "Aurora", "Vector", "Nexus", "Orion", "Volt", "Sigma", "Delta", "Lumen"]

    tipo_by_id = dims["Dim_TipoProjeto"].set_index("ID_TipoProjeto").to_dict("index")
    status_by_name = dims["Dim_Status"].set_index("Status").to_dict("index")
    etapa_by_order = dims["Dim_Etapa"].set_index("Ordem_Etapa").to_dict("index")

    rows = []
    start_base = date(2024, 1, 10)
    for idx in range(1, n + 1):
        tipo_id = random.choices(dims["Dim_TipoProjeto"]["ID_TipoProjeto"], weights=tipos_weights, k=1)[0]
        tipo = tipo_by_id[tipo_id]["Tipo_Projeto"]
        status = random.choices(dims["Dim_Status"]["Status"], weights=status_weights, k=1)[0]
        complexity = weighted_choice(complexities, [18, 47, 27, 8])
        priority = weighted_choice(priorities, [16, 50, 26, 8])
        lead = project_duration_days(tipo, complexity)
        data_entrada = start_base + timedelta(days=random.randint(0, 900))
        if status not in ["Concluido", "Cancelado", "Suspenso", "Nao iniciado"]:
            timing = random.random()
            if timing < 0.48:
                data_entrada = REFERENCE_DATE - timedelta(days=random.randint(0, max(8, lead - 7)))
            elif timing < 0.78:
                data_entrada = REFERENCE_DATE - timedelta(days=random.randint(lead + 1, lead + 85))
            else:
                data_entrada = REFERENCE_DATE - timedelta(days=random.randint(max(10, lead // 2), lead + 20))
        elif status == "Nao iniciado":
            data_entrada = REFERENCE_DATE + timedelta(days=random.randint(-20, 60))
        data_inicio = data_entrada + timedelta(days=random.randint(1, 20)) if status != "Nao iniciado" else None
        data_prazo = data_entrada + timedelta(days=lead)

        if status == "Concluido":
            actual_duration = max(15, int(lead * random.uniform(0.75, 1.22)))
            data_conclusao = data_entrada + timedelta(days=actual_duration)
            etapa_order = 13
            pct = 100
        elif status == "Cancelado":
            data_conclusao = None
            etapa_order = random.randint(1, 7)
            pct = random.choice([0, 10, 20, 35, 50])
        elif status == "Suspenso":
            data_conclusao = None
            etapa_order = random.randint(2, 10)
            pct = min(90, etapa_order * random.randint(6, 8))
        elif status == "Nao iniciado":
            data_conclusao = None
            etapa_order = 1
            pct = 0
        else:
            data_conclusao = None
            status_to_stage = {
                "Em analise": (2, 3), "Em desenvolvimento": (4, 5),
                "Em aprovacao": (6, 9), "Aguardando cliente": (3, 8),
                "Aguardando concessionaria": (6, 9), "Em execucao": (10, 11),
                "Em comissionamento": (12, 12),
            }
            lo, hi = status_to_stage[status]
            etapa_order = random.randint(lo, hi)
            pct = min(96, max(5, int((etapa_order - 1) / 12 * 100 + random.randint(-8, 8))))

        potencia = power_for_type(tipo)
        level = tipo_by_id[tipo_id]["Nivel_Tensao_Padrao"]
        if potencia > 75 and tipo not in ["Entrada de Energia BT", "Microgeracao Distribuida"]:
            level = "MT"
        tipo_conexao = "Grupo A" if level == "MT" else "Grupo B"
        tipo_sistema = tipo_by_id[tipo_id]["Categoria_Engenharia"]
        value_per_kw = random.uniform(2300, 5200)
        service_base = {
            "Protecao e estudos": 1.2, "Energia e operacoes": 0.75, "Automacao": 1.45,
            "Infraestrutura eletrica": 1.55, "Mobilidade eletrica": 1.3,
            "Entrada de energia": 0.9, "Geracao distribuida": 1.0,
        }[tipo_sistema]
        valor = brl_round(max(8500, potencia * value_per_kw * service_base))
        custo = brl_round(valor * random.uniform(0.54, 0.78))
        horas_estimadas = int(max(18, lead * random.uniform(0.6, 1.7) + potencia / 18))
        horas_realizadas = int(horas_estimadas * (pct / 100) * random.uniform(0.85, 1.25))
        if status == "Concluido":
            horas_realizadas = int(horas_estimadas * random.uniform(0.88, 1.22))

        overdue = status not in ["Concluido", "Cancelado", "Suspenso"] and data_prazo < REFERENCE_DATE
        base_simple = status_by_name[status]["Status_Simplificado_Base"]
        status_simplificado = "ATRASADO" if overdue else base_simple
        days_late = max(0, (REFERENCE_DATE - data_prazo).days) if overdue else 0
        stopped_days = 0
        if status in ["Aguardando cliente", "Aguardando concessionaria", "Suspenso"] or overdue:
            stopped_days = random.randint(5, 55) + min(days_late, 40)

        has_issue_probability = 0.22 + (0.18 if overdue else 0) + (0.12 if status in ["Aguardando cliente", "Aguardando concessionaria"] else 0)
        possui_pendencia = random.random() < has_issue_probability and status not in ["Concluido", "Cancelado"]
        qtd_pendencias = random.randint(1, 4) if possui_pendencia else 0

        cliente_id = random.choice(dims["Dim_Cliente"]["ID_Cliente"].tolist())
        resp_id = random.choice(dims["Dim_Responsavel"]["ID_Responsavel"].tolist())
        conc_id = random.choices(dims["Dim_Concessionaria"]["ID_Concessionaria"], weights=conc_weights, k=1)[0]
        estado = random.choices(states, weights=state_weights, k=1)[0]
        loc_id = dims["Dim_Localidade"].loc[dims["Dim_Localidade"]["Estado"] == estado, "ID_Localidade"].iloc[0]

        rows.append({
            "ID_Projeto": f"PROJ{idx:04d}",
            "Nome_Projeto": f"Projeto {random.choice(project_prefixes)} {idx:03d}",
            "ID_Cliente": cliente_id,
            "ID_Responsavel": resp_id,
            "ID_Concessionaria": conc_id,
            "ID_TipoProjeto": tipo_id,
            "ID_Status": status_by_name[status]["ID_Status"],
            "ID_Etapa": etapa_by_order[etapa_order]["ID_Etapa"],
            "ID_Localidade": loc_id,
            "Data_Entrada": data_entrada,
            "Data_Inicio": data_inicio,
            "Data_Prazo": data_prazo,
            "Data_Conclusao": data_conclusao,
            "Potencia_kW": potencia,
            "Valor_Projeto": valor,
            "Custo_Estimado": custo,
            "Margem_Estimada": brl_round(valor - custo),
            "Horas_Estimadas": horas_estimadas,
            "Horas_Realizadas": horas_realizadas,
            "Percentual_Conclusao": pct,
            "Prioridade": priority,
            "Complexidade": complexity,
            "Nivel_Tensao": level,
            "Tipo_Conexao": tipo_conexao,
            "Tipo_Sistema": tipo_sistema,
            "Possui_Pendencia": possui_pendencia,
            "Quantidade_Pendencias": qtd_pendencias,
            "Dias_Parado": stopped_days,
            "Status_Simplificado": status_simplificado,
            "Dias_Ate_Prazo": (data_prazo - REFERENCE_DATE).days if status not in ["Concluido", "Cancelado"] else None,
            "Dias_Atraso": days_late,
            "Faixa_Prazo": prazo_bucket((data_prazo - REFERENCE_DATE).days, status_simplificado),
            "Criticidade": criticidade(status_simplificado, priority, days_late, qtd_pendencias),
            "Observacao_Status": status_note(status, status_simplificado, qtd_pendencias),
        })
    return pd.DataFrame(rows)


def prazo_bucket(days_until: int | None, simplified: str) -> str:
    if simplified == "CONCLUIDO":
        return "Concluido"
    if simplified == "CANCELADO":
        return "Cancelado"
    if simplified == "SUSPENSO":
        return "Suspenso"
    if days_until is None:
        return "n.a."
    if days_until < 0:
        return "Vencido"
    if days_until <= 7:
        return "0-7 dias"
    if days_until <= 15:
        return "8-15 dias"
    if days_until <= 30:
        return "16-30 dias"
    return "Mais de 30 dias"


def criticidade(simplified: str, priority: str, days_late: int, issues: int) -> str:
    if simplified == "CANCELADO":
        return "Encerrado sem acao"
    if simplified == "CONCLUIDO":
        return "Baixa"
    score = 0
    score += {"Baixa": 0, "Media": 1, "Alta": 2, "Critica": 3}[priority]
    score += 3 if days_late > 30 else 2 if days_late > 10 else 1 if days_late > 0 else 0
    score += min(2, issues)
    if simplified == "SUSPENSO":
        score += 2
    if score >= 6:
        return "Critica"
    if score >= 4:
        return "Alta"
    if score >= 2:
        return "Media"
    return "Baixa"


def status_note(status: str, simplified: str, issues: int) -> str:
    if simplified == "ATRASADO":
        return "Prazo vencido com projeto ainda aberto"
    if status == "Aguardando concessionaria":
        return "Dependencia externa da distribuidora"
    if status == "Aguardando cliente":
        return "Aguardando documentos ou decisao do cliente"
    if issues:
        return "Projeto com pendencias registradas"
    if status == "Concluido":
        return "Projeto encerrado"
    return "Fluxo regular"


def create_pending(projects: pd.DataFrame) -> pd.DataFrame:
    categories = [
        "Documento cliente", "Informacao tecnica", "Aprovacao interna",
        "Concessionaria", "Alteracao de projeto", "Vistoria", "Pagamento",
        "Equipamento", "Execucao", "Outros"
    ]
    rows = []
    pend_id = 1
    for _, row in projects.iterrows():
        for _ in range(int(row["Quantidade_Pendencias"])):
            opened = pd.to_datetime(row["Data_Entrada"]).date() + timedelta(days=random.randint(5, 95))
            closed = None
            status = weighted_choice(["Aberta", "Em tratativa", "Resolvida"], [45, 35, 20])
            if status == "Resolvida":
                closed = opened + timedelta(days=random.randint(2, 25))
            rows.append({
                "ID_Pendencia": f"PEND{pend_id:04d}",
                "ID_Projeto": row["ID_Projeto"],
                "Categoria_Pendencia": weighted_choice(categories, [17, 13, 10, 18, 11, 8, 7, 6, 6, 4]),
                "Descricao_Pendencia": "Pendencia sintetica para demonstracao de gestao operacional",
                "Data_Abertura": opened,
                "Data_Resolucao": closed,
                "Status_Pendencia": status,
                "Dias_Em_Aberto": max(0, ((REFERENCE_DATE if closed is None else closed) - opened).days),
                "Impacta_Prazo": random.random() < 0.62,
            })
            pend_id += 1
    return pd.DataFrame(rows)


def create_status_history(projects: pd.DataFrame, dims: dict[str, pd.DataFrame]) -> pd.DataFrame:
    status_lookup = dims["Dim_Status"].set_index("ID_Status")["Status"].to_dict()
    rows = []
    hist_id = 1
    sequence = ["Nao iniciado", "Em analise", "Em desenvolvimento", "Em aprovacao", "Aguardando concessionaria", "Em execucao", "Em comissionamento", "Concluido"]
    for _, p in projects.iterrows():
        current = status_lookup[p["ID_Status"]]
        if current in sequence:
            until = sequence.index(current) + 1
            chosen = sequence[:until]
        elif current == "Suspenso":
            chosen = sequence[:random.randint(2, 5)] + ["Suspenso"]
        elif current == "Cancelado":
            chosen = sequence[:random.randint(1, 4)] + ["Cancelado"]
        else:
            chosen = ["Nao iniciado", current]
        start = pd.to_datetime(p["Data_Entrada"]).date()
        for order, st in enumerate(chosen, start=1):
            rows.append({
                "ID_Historico": f"HIST{hist_id:05d}",
                "ID_Projeto": p["ID_Projeto"],
                "Status": st,
                "Data_Status": start + timedelta(days=(order - 1) * random.randint(9, 22)),
                "Ordem_Status": order,
            })
            hist_id += 1
    return pd.DataFrame(rows)


def create_hours(projects: pd.DataFrame) -> pd.DataFrame:
    rows = []
    hid = 1
    for _, p in projects.iterrows():
        total = int(p["Horas_Realizadas"])
        if total <= 0:
            continue
        months = max(1, min(8, total // 25 + 1))
        remaining = total
        start = pd.to_datetime(p["Data_Inicio"] if pd.notna(p["Data_Inicio"]) else p["Data_Entrada"]).date()
        for m in range(months):
            spent = remaining if m == months - 1 else max(1, int(total / months * random.uniform(0.75, 1.25)))
            remaining = max(0, remaining - spent)
            rows.append({
                "ID_Horas": f"HORA{hid:05d}",
                "ID_Projeto": p["ID_Projeto"],
                "ID_Responsavel": p["ID_Responsavel"],
                "Mes_Referencia": (start + timedelta(days=30 * m)).replace(day=1),
                "Horas": spent,
                "Atividade": weighted_choice(["Analise", "Projeto", "Revisao", "Protocolo", "Campo", "Gestao"], [18, 34, 14, 10, 14, 10]),
            })
            hid += 1
            if remaining <= 0:
                break
    return pd.DataFrame(rows)


def create_faturamento(projects: pd.DataFrame) -> pd.DataFrame:
    rows = []
    fid = 1
    for _, p in projects.iterrows():
        milestones = [("Entrada", 0.25), ("Projeto", 0.35), ("Aprovacao", 0.25), ("Encerramento", 0.15)]
        start = pd.to_datetime(p["Data_Entrada"]).date()
        for idx, (marco, pct) in enumerate(milestones):
            if p["Percentual_Conclusao"] < [5, 35, 70, 100][idx]:
                continue
            rows.append({
                "ID_Faturamento": f"FAT{fid:05d}",
                "ID_Projeto": p["ID_Projeto"],
                "Marco": marco,
                "Data_Faturamento": start + timedelta(days=idx * 35 + random.randint(0, 10)),
                "Valor_Faturado": brl_round(p["Valor_Projeto"] * pct),
            })
            fid += 1
    return pd.DataFrame(rows)


def save_outputs(tables: dict[str, pd.DataFrame]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    localized_tables = {name: df.replace(LOCALIZE_VALUES) for name, df in tables.items()}
    for name, df in tables.items():
        localized_tables[name].to_csv(DATA_DIR / f"{name}.csv", index=False, encoding="utf-8-sig")

    xlsx_path = ROOT / "dashboard_data.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        for name, df in localized_tables.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
    format_workbook(xlsx_path)


def format_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="12372A")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        max_row = ws.max_row
        max_col = ws.max_column
        if max_row > 1 and max_col > 1:
            ref = f"A1:{ws.cell(max_row, max_col).coordinate}"
            table = Table(displayName=f"tbl_{ws.title.replace(' ', '_')[:24]}", ref=ref)
            style = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False)
            table.tableStyleInfo = style
            ws.add_table(table)
        for col_cells in ws.columns:
            letter = col_cells[0].column_letter
            width = min(32, max(10, max(len(str(c.value)) if c.value is not None else 0 for c in col_cells[:80]) + 2))
            ws.column_dimensions[letter].width = width
    wb.save(path)


def main() -> None:
    random.seed(SEED)
    np.random.seed(SEED)
    dims = build_dimensions()
    projects = create_projects(dims, 200)
    pending = create_pending(projects)
    history = create_status_history(projects, dims)
    hours = create_hours(projects)
    billing = create_faturamento(projects)
    tables = {**dims, "Fato_Projetos": projects, "Fato_Pendencias": pending, "Fato_HistoricoStatus": history, "Fato_Horas": hours, "Fato_Faturamento": billing}
    ordered = {
        "Fato_Projetos": tables["Fato_Projetos"],
        "Dim_Cliente": tables["Dim_Cliente"],
        "Dim_Responsavel": tables["Dim_Responsavel"],
        "Dim_Concessionaria": tables["Dim_Concessionaria"],
        "Dim_TipoProjeto": tables["Dim_TipoProjeto"],
        "Dim_Status": tables["Dim_Status"],
        "Dim_Etapa": tables["Dim_Etapa"],
        "Dim_Localidade": tables["Dim_Localidade"],
        "Dim_Calendario": tables["Dim_Calendario"],
        "Fato_Pendencias": tables["Fato_Pendencias"],
        "Fato_HistoricoStatus": tables["Fato_HistoricoStatus"],
        "Fato_Horas": tables["Fato_Horas"],
        "Fato_Faturamento": tables["Fato_Faturamento"],
    }
    save_outputs(ordered)
    print(f"Gerados {len(projects)} projetos, {len(pending)} pendências e {len(history)} registros de histórico de status.")
    print(f"Data de referência: {REFERENCE_DATE.isoformat()}")


if __name__ == "__main__":
    main()
