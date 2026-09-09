from __future__ import annotations

import hashlib
import json
import shutil
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
DOCS = ROOT / "docs"
POWERBI = ROOT / "powerbi"
MOCKUPS = ROOT / "assets" / "mockups"
SCREENSHOTS = ROOT / "assets" / "screenshots"
ARCH = ROOT / "assets" / "architecture"
LINKEDIN = ROOT / "linkedin"
WEB = ROOT / "web"
OUTPUTS = ROOT.parent / "outputs"
REFERENCE_DATE = pd.Timestamp("2026-09-08")
REPO_URL = "https://github.com/renovera1/electrical-project-management-dashboard"
DASHBOARD_URL = "https://renovera1.github.io/electrical-project-management-dashboard/"

COLORS = {
    "green": "12372A",
    "green2": "1F5B43",
    "mint": "DDEBE4",
    "bg": "F6F8F7",
    "line": "DDE4E0",
    "text": "20252B",
    "muted": "69737D",
    "gold": "C8A24A",
    "amber": "B7791F",
    "red": "B42318",
    "blue": "2F6B8F",
    "white": "FFFFFF",
}


def hx(name: str) -> str:
    return "#" + COLORS[name]


def br_int(value: float) -> str:
    return f"{int(round(float(value))):,}".replace(",", ".")


def br_dec(value: float, places: int = 1) -> str:
    raw = f"{float(value):,.{places}f}"
    return raw.replace(",", "X").replace(".", ",").replace("X", ".")


def br_money(value: float) -> str:
    value = float(value)
    if abs(value) >= 1_000_000:
        return f"R$ {br_dec(value / 1_000_000, 1)} mi"
    if abs(value) >= 1_000:
        return f"R$ {br_dec(value / 1_000, 0)} mil"
    return "R$ " + br_dec(value, 2)


def br_pct(value: float) -> str:
    return f"{br_dec(float(value) * 100, 1)}%"


def ensure_dirs() -> None:
    for path in [DOCS, POWERBI, MOCKUPS, SCREENSHOTS, ARCH, LINKEDIN, WEB, OUTPUTS]:
        path.mkdir(parents=True, exist_ok=True)


def read_tables() -> dict[str, pd.DataFrame]:
    return {path.stem: pd.read_csv(path, encoding="utf-8-sig") for path in DATA.glob("*.csv")}


def join_projects(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    df = tables["Fato_Projetos"].copy()
    for name, key in [
        ("Dim_Cliente", "ID_Cliente"),
        ("Dim_Responsavel", "ID_Responsavel"),
        ("Dim_Concessionaria", "ID_Concessionaria"),
        ("Dim_TipoProjeto", "ID_TipoProjeto"),
        ("Dim_Status", "ID_Status"),
        ("Dim_Etapa", "ID_Etapa"),
        ("Dim_Localidade", "ID_Localidade"),
    ]:
        df = df.merge(tables[name], on=key, how="left")
    for col in ["Data_Entrada", "Data_Inicio", "Data_Prazo", "Data_Conclusao"]:
        df[col] = pd.to_datetime(df[col], errors="coerce")
    return df


def metrics(df: pd.DataFrame, pend: pd.DataFrame) -> dict[str, object]:
    concluded = df["Status_Simplificado"].eq("CONCLUÍDO")
    active = ~df["Status_Simplificado"].isin(["CONCLUÍDO", "CANCELADO", "SUSPENSO"])
    overdue = df["Status_Simplificado"].eq("ATRASADO")
    on_time = concluded & (df["Data_Conclusao"] <= df["Data_Prazo"])
    value = float(df["Valor_Projeto"].sum())
    cost = float(df["Custo_Estimado"].sum())
    estimated_hours = int(df["Horas_Estimadas"].sum())
    actual_hours = int(df["Horas_Realizadas"].sum())
    return {
        "data_referencia": REFERENCE_DATE.strftime("%d/%m/%Y"),
        "total_projetos": int(df["ID_Projeto"].nunique()),
        "projetos_ativos": int(active.sum()),
        "projetos_concluidos": int(concluded.sum()),
        "projetos_atrasados": int(overdue.sum()),
        "projetos_suspensos": int(df["Status_Simplificado"].eq("SUSPENSO").sum()),
        "projetos_cancelados": int(df["Status_Simplificado"].eq("CANCELADO").sum()),
        "potencia_total_kw": float(df["Potencia_kW"].sum()),
        "potencia_total_mw": float(df["Potencia_kW"].sum() / 1000),
        "valor_carteira": value,
        "valor_ativo": float(df.loc[active, "Valor_Projeto"].sum()),
        "custo_estimado": cost,
        "margem_estimada": value - cost,
        "ticket_medio": float(df["Valor_Projeto"].mean()),
        "prazo_medio_dias": float((df["Data_Prazo"] - df["Data_Entrada"]).dt.days.mean()),
        "taxa_conclusao": float(concluded.mean()),
        "taxa_atraso": float(overdue.mean()),
        "taxa_concluido_prazo": float(on_time.sum() / concluded.sum()) if concluded.sum() else 0.0,
        "total_pendencias": int(len(pend)),
        "projetos_com_pendencia": int(df["Possui_Pendencia"].sum()),
        "dias_parado_medio": float(df["Dias_Parado"].mean()),
        "horas_estimadas": estimated_hours,
        "horas_realizadas": actual_hours,
        "desvio_horas": actual_hours - estimated_hours,
        "eficiencia_horas": float(estimated_hours / actual_hours) if actual_hours else 0.0,
    }


def series_dict(series: pd.Series, limit: int = 12, decimals: int | None = None) -> dict[str, float | int]:
    out: dict[str, float | int] = {}
    for key, value in series.sort_values(ascending=False).head(limit).items():
        out[str(key)] = round(float(value), decimals) if decimals is not None else int(value)
    return out


def insights(df: pd.DataFrame, pend: pd.DataFrame, k: dict[str, object]) -> list[str]:
    top_conc = df.groupby("Concessionaria")["ID_Projeto"].nunique().sort_values(ascending=False).head(1)
    top_state = df.groupby("Estado")["ID_Projeto"].nunique().sort_values(ascending=False).head(1)
    overdue_type = df[df["Status_Simplificado"].eq("ATRASADO")].groupby("Tipo_Projeto")["ID_Projeto"].nunique().sort_values(ascending=False).head(1)
    workload = df[~df["Status_Simplificado"].isin(["CONCLUÍDO", "CANCELADO"])].groupby("Responsavel")["ID_Projeto"].nunique().sort_values(ascending=False).head(1)
    pend_cat = pend.groupby("Categoria_Pendencia")["ID_Pendencia"].count().sort_values(ascending=False).head(1)
    power_share = df.sort_values("Potencia_kW", ascending=False).head(10)["Potencia_kW"].sum() / df["Potencia_kW"].sum()
    stage = df.groupby("Macro_Etapa")["Dias_Parado"].mean().sort_values(ascending=False).head(1)
    return [
        f"{top_conc.index[0]} concentra {int(top_conc.iloc[0])} projetos, a maior quantidade por concessionária.",
        f"{top_state.index[0]} é o estado com maior volume, com {int(top_state.iloc[0])} projetos na carteira.",
        f"{overdue_type.index[0]} aparece como o tipo com mais atrasos, com {int(overdue_type.iloc[0])} projetos atrasados.",
        f"{workload.index[0]} tem a maior carga operacional aberta, com {int(workload.iloc[0])} projetos não encerrados.",
        f"Os 10 maiores projetos concentram {br_pct(power_share)} da potência total.",
        f"A macroetapa {stage.index[0]} tem o maior tempo médio parado, com {br_dec(stage.iloc[0], 1)} dias.",
        f"{pend_cat.index[0]} é a categoria de pendência mais frequente, com {int(pend_cat.iloc[0])} registros.",
        f"A carteira soma {br_dec(k['potencia_total_mw'], 1)} MW e {br_money(k['valor_carteira'])}, conectando risco técnico e impacto financeiro.",
    ]


def build_summary(df: pd.DataFrame, pend: pd.DataFrame, tables: dict[str, pd.DataFrame]) -> dict[str, object]:
    k = metrics(df, pend)
    pend_counts = pend.groupby("ID_Projeto").size().rename("Pendencias_Registradas")
    enriched = df.merge(pend_counts, on="ID_Projeto", how="left")
    enriched["Pendencias_Registradas"] = enriched["Pendencias_Registradas"].fillna(0).astype(int)
    cols = [
        "ID_Projeto", "Nome_Projeto", "Cliente", "Segmento", "Responsavel", "Cargo", "Concessionaria",
        "Tipo_Projeto", "Categoria_Engenharia", "Status", "Status_Simplificado", "Etapa", "Macro_Etapa",
        "Estado", "Municipio", "Regiao", "Data_Entrada", "Data_Inicio", "Data_Prazo", "Data_Conclusao",
        "Potencia_kW", "Valor_Projeto", "Custo_Estimado", "Margem_Estimada", "Horas_Estimadas",
        "Horas_Realizadas", "Percentual_Conclusao", "Prioridade", "Complexidade", "Nivel_Tensao",
        "Tipo_Conexao", "Tipo_Sistema", "Quantidade_Pendencias", "Pendencias_Registradas",
        "Dias_Parado", "Dias_Ate_Prazo", "Dias_Atraso", "Faixa_Prazo", "Criticidade", "Observacao_Status",
    ]
    records = enriched[cols].copy()
    for col in ["Data_Entrada", "Data_Inicio", "Data_Prazo", "Data_Conclusao"]:
        records[col] = records[col].dt.strftime("%Y-%m-%d").fillna("")
    return {
        "metadata": {
            "titulo": "Dashboard de Gestão de Projetos Elétricos",
            "subtitulo": "Dashboard de gestão de projetos de engenharia elétrica",
            "data_referencia": k["data_referencia"],
            "origem": "Dados sintéticos gerados com Python.",
        },
        "metricas": k,
        "metricas_formatadas": {
            "total_projetos": br_int(k["total_projetos"]),
            "projetos_ativos": br_int(k["projetos_ativos"]),
            "projetos_concluidos": br_int(k["projetos_concluidos"]),
            "projetos_atrasados": br_int(k["projetos_atrasados"]),
            "potencia_total_mw": f"{br_dec(k['potencia_total_mw'], 1)} MW",
            "valor_carteira": br_money(k["valor_carteira"]),
            "valor_ativo": br_money(k["valor_ativo"]),
            "margem_estimada": br_money(k["margem_estimada"]),
            "ticket_medio": br_money(k["ticket_medio"]),
            "taxa_conclusao": br_pct(k["taxa_conclusao"]),
            "taxa_atraso": br_pct(k["taxa_atraso"]),
            "taxa_concluido_prazo": br_pct(k["taxa_concluido_prazo"]),
            "total_pendencias": br_int(k["total_pendencias"]),
            "horas_estimadas": br_int(k["horas_estimadas"]),
            "horas_realizadas": br_int(k["horas_realizadas"]),
            "eficiencia_horas": br_dec(k["eficiencia_horas"], 2),
        },
        "graficos": {
            "status": series_dict(df.groupby("Status_Simplificado")["ID_Projeto"].nunique()),
            "tipo_projeto": series_dict(df.groupby("Tipo_Projeto")["ID_Projeto"].nunique()),
            "estado": series_dict(df.groupby("Estado")["ID_Projeto"].nunique()),
            "concessionaria": series_dict(df.groupby("Concessionaria")["ID_Projeto"].nunique()),
            "responsavel": series_dict(df.groupby("Responsavel")["ID_Projeto"].nunique()),
            "etapa": series_dict(df.groupby("Macro_Etapa")["ID_Projeto"].nunique()),
            "prazo": series_dict(df.groupby("Faixa_Prazo")["ID_Projeto"].nunique()),
            "prioridade": series_dict(df.groupby("Prioridade")["ID_Projeto"].nunique()),
            "categoria_pendencia": series_dict(pend.groupby("Categoria_Pendencia")["ID_Pendencia"].count()),
            "potencia_categoria": series_dict(df.groupby("Categoria_Engenharia")["Potencia_kW"].sum(), decimals=1),
            "valor_tipo": series_dict(df.groupby("Tipo_Projeto")["Valor_Projeto"].sum(), decimals=2),
        },
        "projetos": records.to_dict("records"),
        "pendencias": pend.to_dict("records"),
        "insights": insights(df, pend, k),
        "filtros": {
            "anos": sorted({int(d.year) for d in df["Data_Entrada"].dropna()}),
            "estados": sorted(df["Estado"].dropna().unique().tolist()),
            "concessionarias": sorted(df["Concessionaria"].dropna().unique().tolist()),
            "tipos": sorted(df["Tipo_Projeto"].dropna().unique().tolist()),
            "responsaveis": sorted(df["Responsavel"].dropna().unique().tolist()),
            "status": sorted(df["Status_Simplificado"].dropna().unique().tolist()),
            "etapas": sorted(df["Macro_Etapa"].dropna().unique().tolist()),
            "prioridades": sorted(df["Prioridade"].dropna().unique().tolist()),
        },
        "tabelas": {name: {"linhas": len(table), "colunas": list(table.columns)} for name, table in tables.items()},
    }


def write_json(summary: dict[str, object]) -> None:
    payload = json.dumps(summary, ensure_ascii=False, indent=2)
    (OUTPUTS / "resumo_portfolio.json").write_text(payload, encoding="utf-8")
    (WEB / "portfolio_summary.json").write_text(payload, encoding="utf-8")
    (WEB / "portfolio_summary.js").write_text("window.PORTFOLIO_SUMMARY = " + payload + ";\n", encoding="utf-8")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for candidate in [
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            pass
    return ImageFont.load_default()


def draw_card(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], title: str, value: str, color: str) -> None:
    draw.rounded_rectangle(box, radius=8, fill=hx("white"), outline=hx("line"))
    draw.rectangle((box[0], box[1], box[0] + 6, box[3]), fill=hx(color))
    draw.text((box[0] + 18, box[1] + 16), title, fill=hx("muted"), font=font(16))
    draw.text((box[0] + 18, box[1] + 45), value, fill=hx("text"), font=font(26, True))


def draw_bars(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], title: str, data: dict[str, float | int], color: str) -> None:
    draw.rounded_rectangle(box, radius=8, fill=hx("white"), outline=hx("line"))
    x1, y1, x2, y2 = box
    draw.text((x1 + 18, y1 + 14), title, fill=hx("text"), font=font(18, True))
    max_value = max([float(v) for v in data.values()] or [1])
    y = y1 + 54
    for label, value in list(data.items())[:8]:
        if y > y2 - 30:
            break
        width = int((x2 - x1 - 220) * float(value) / max_value)
        draw.text((x1 + 18, y), label[:25], fill=hx("muted"), font=font(13))
        draw.rounded_rectangle((x1 + 210, y + 3, x1 + 210 + width, y + 19), radius=4, fill=hx(color))
        draw.text((x2 - 70, y), br_dec(value, 0), fill=hx("text"), font=font(13, True))
        y += 31


def make_mockups(summary: dict[str, object]) -> None:
    m, charts = summary["metricas_formatadas"], summary["graficos"]
    pages = [
        ("executive_overview.png", "Visão Executiva", [("Projetos", m["total_projetos"]), ("Ativos", m["projetos_ativos"]), ("Atrasados", m["projetos_atrasados"]), ("Potência", m["potencia_total_mw"]), ("Carteira", m["valor_carteira"])], "Projetos por status", charts["status"], "Projetos por tipo", charts["tipo_projeto"]),
        ("deadlines.png", "Gestão de Prazos", [("Atrasados", m["projetos_atrasados"]), ("Taxa atraso", m["taxa_atraso"]), ("No prazo", m["taxa_concluido_prazo"]), ("Pendências", m["total_pendencias"])], "Faixas de prazo", charts["prazo"], "Projetos por concessionária", charts["concessionaria"]),
        ("operations.png", "Operações", [("Horas est.", m["horas_estimadas"]), ("Horas real.", m["horas_realizadas"]), ("Eficiência", m["eficiencia_horas"]), ("Pendências", m["total_pendencias"])], "Projetos por etapa", charts["etapa"], "Carga por responsável", charts["responsavel"]),
        ("engineering.png", "Engenharia", [("Potência", m["potencia_total_mw"]), ("Projetos", m["total_projetos"]), ("Ativos", m["projetos_ativos"]), ("Concluídos", m["projetos_concluidos"])], "Potência por categoria", charts["potencia_categoria"], "Projetos por estado", charts["estado"]),
        ("financial.png", "Financeiro", [("Carteira", m["valor_carteira"]), ("Ativo", m["valor_ativo"]), ("Margem", m["margem_estimada"]), ("Ticket", m["ticket_medio"])], "Valor por tipo", charts["valor_tipo"], "Projetos por status", charts["status"]),
        ("issues.png", "Pendências e Gargalos", [("Pendências", m["total_pendencias"]), ("Projetos", m["total_projetos"]), ("Atrasados", m["projetos_atrasados"]), ("Taxa atraso", m["taxa_atraso"])], "Pendências por categoria", charts["categoria_pendencia"], "Projetos por etapa", charts["etapa"]),
    ]
    for filename, title, cards, left_title, left, right_title, right in pages:
        img = Image.new("RGB", (1440, 900), hx("bg"))
        draw = ImageDraw.Draw(img)
        draw.rectangle((0, 0, 1440, 78), fill=hx("green"))
        draw.text((34, 22), title, fill=hx("white"), font=font(28, True))
        draw.text((1010, 28), "Dashboard de Gestão de Projetos Elétricos", fill="#D7E5DD", font=font(15))
        for idx, (label, value) in enumerate(cards):
            draw_card(draw, (34 + idx * 270, 110, 284 + idx * 270, 205), label, str(value), "red" if "Atras" in label else "green2")
        draw_bars(draw, (34, 245, 690, 775), left_title, left, "green2")
        draw_bars(draw, (730, 245, 1395, 775), right_title, right, "blue")
        img.save(MOCKUPS / filename)
        shutil.copyfile(MOCKUPS / filename, SCREENSHOTS / filename)
    detail = Image.new("RGB", (1440, 900), hx("bg"))
    draw = ImageDraw.Draw(detail)
    draw.rectangle((0, 0, 1440, 78), fill=hx("green"))
    draw.text((34, 22), "Detalhamento do Projeto", fill=hx("white"), font=font(28, True))
    project = sorted(summary["projetos"], key=lambda p: (p["Criticidade"], p["Dias_Atraso"]), reverse=True)[0]
    items = [
        ("Projeto", project["Nome_Projeto"]), ("Cliente", project["Cliente"]), ("Município", f"{project['Municipio']} / {project['Estado']}"),
        ("Concessionária", project["Concessionaria"]), ("Responsável", project["Responsavel"]), ("Status", project["Status_Simplificado"]),
        ("Tipo", project["Tipo_Projeto"]), ("Etapa", project["Etapa"]), ("Potência", f"{br_dec(project['Potencia_kW'], 1)} kW"),
        ("Valor", br_money(project["Valor_Projeto"])), ("Prazo", project["Data_Prazo"]), ("Pendências", str(project["Quantidade_Pendencias"])),
    ]
    for idx, (label, value) in enumerate(items):
        x, y = 55 + (idx % 3) * 440, 130 + (idx // 3) * 120
        draw_card(draw, (x, y, x + 390, y + 92), label, str(value), "red" if label == "Status" else "green2")
    detail.save(MOCKUPS / "project_detail.png")
    shutil.copyfile(MOCKUPS / "project_detail.png", SCREENSHOTS / "project_detail.png")
    make_linkedin_images(m)


def make_linkedin_images(m: dict[str, str]) -> None:
    cover = Image.new("RGB", (1200, 627), hx("green"))
    draw = ImageDraw.Draw(cover)
    draw.rectangle((0, 470, 1200, 627), fill=hx("gold"))
    draw.text((70, 80), "Dashboard de Gestão\nde Projetos Elétricos", fill=hx("white"), font=font(58, True), spacing=8)
    draw.text((70, 270), f"{m['total_projetos']} projetos | {m['potencia_total_mw']} | {m['valor_carteira']}", fill="#EAF3EE", font=font(28))
    draw.text((70, 515), "Power BI | Python | DAX | Gestão de Engenharia", fill=hx("text"), font=font(30, True))
    cover.save(LINKEDIN / "capa_linkedin_1200x627.png")
    square = Image.new("RGB", (1080, 1080), hx("bg"))
    draw = ImageDraw.Draw(square)
    draw.rectangle((0, 0, 1080, 210), fill=hx("green"))
    draw.text((60, 58), "Portfólio Power BI\nEngenharia Elétrica", fill=hx("white"), font=font(50, True), spacing=4)
    for i, (label, value) in enumerate([("Projetos", m["total_projetos"]), ("Potência", m["potencia_total_mw"]), ("Carteira", m["valor_carteira"]), ("Atrasados", m["projetos_atrasados"])]):
        x, y = 70 + (i % 2) * 480, 300 + (i // 2) * 230
        draw_card(draw, (x, y, x + 410, y + 150), label, value, "red" if label == "Atrasados" else "green2")
    draw.text((70, 850), "Dataset sintético, modelo estrela, DAX,\nPower Query, Excel e demo web filtrável.", fill=hx("text"), font=font(31), spacing=8)
    square.save(LINKEDIN / "arte_linkedin_1080x1080.png")


def make_architecture() -> None:
    img = Image.new("RGB", (1400, 850), hx("white"))
    draw = ImageDraw.Draw(img)
    draw.text((395, 35), "Modelo de Dados em Estrela", fill=hx("text"), font=font(34, True))
    draw.rounded_rectangle((540, 310, 860, 525), radius=10, fill=hx("green"))
    draw.text((608, 397), "Fato_Projetos", fill=hx("white"), font=font(26, True))
    dims = ["Dim_Cliente", "Dim_Responsavel", "Dim_Concessionaria", "Dim_TipoProjeto", "Dim_Status", "Dim_Etapa", "Dim_Localidade", "Dim_Calendario", "Fato_Pendencias", "Fato_Horas", "Fato_Faturamento"]
    positions = [(70, 120), (410, 120), (910, 120), (1120, 320), (930, 650), (580, 650), (230, 650), (60, 350), (1055, 520), (270, 320), (620, 120)]
    for name, (x, y) in zip(dims, positions):
        draw.line((x + 115, y + 72, 700, 418), fill=hx("muted"), width=2)
        draw.rounded_rectangle((x, y, x + 235, y + 72), radius=8, fill=hx("bg"), outline=hx("line"))
        draw.text((x + 17, y + 24), name, fill=hx("text"), font=font(16, True))
    img.save(ARCH / "data_model.png")


def style_header(cell) -> None:
    cell.fill = PatternFill("solid", fgColor=COLORS["green"])
    cell.font = Font(name="Arial", bold=True, color=COLORS["white"])
    cell.alignment = Alignment(horizontal="center", vertical="center")


def write_small_table(ws, start: str, title: str, data: dict[str, float | int], limit: int = 10) -> None:
    cell = ws[start]
    r, c = cell.row, cell.column
    ws.cell(r, c).value = title
    style_header(ws.cell(r, c))
    ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)
    ws.cell(r + 1, c).value = "Categoria"
    ws.cell(r + 1, c + 1).value = "Valor"
    style_header(ws.cell(r + 1, c))
    style_header(ws.cell(r + 1, c + 1))
    for idx, (label, value) in enumerate(list(data.items())[:limit], start=r + 2):
        ws.cell(idx, c).value = label
        ws.cell(idx, c + 1).value = value
        ws.cell(idx, c + 1).number_format = '#,##0.0' if isinstance(value, float) else '#,##0'


def setup_dashboard(ws, summary: dict[str, object]) -> None:
    ws.sheet_properties.tabColor = COLORS["green"]
    ws["B2"] = "Dashboard de Gestão de Projetos Elétricos"
    ws["B2"].font = Font(name="Arial", size=16, bold=True, color=COLORS["text"])
    ws["B3"] = f"Data de referência: {summary['metricas']['data_referencia']} | Dados sintéticos"
    ws["B3"].font = Font(name="Arial", size=10, italic=True, color=COLORS["muted"])
    cards = [
        ("Total de projetos", summary["metricas"]["total_projetos"], "#,##0"),
        ("Projetos ativos", summary["metricas"]["projetos_ativos"], "#,##0"),
        ("Projetos atrasados", summary["metricas"]["projetos_atrasados"], "#,##0"),
        ("Potência total (MW)", summary["metricas"]["potencia_total_mw"], "#,##0.0"),
        ("Valor da carteira", summary["metricas"]["valor_carteira"], '"R$"#,##0'),
        ("Pendências", summary["metricas"]["total_pendencias"], "#,##0"),
    ]
    for idx, (label, value, fmt) in enumerate(cards):
        col = 2 + idx * 2
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
        ws.merge_cells(start_row=6, start_column=col, end_row=6, end_column=col + 1)
        ws.cell(5, col).value = label
        ws.cell(6, col).value = value
        ws.cell(5, col).font = Font(name="Arial", size=9, color=COLORS["muted"])
        ws.cell(6, col).font = Font(name="Arial", size=14, bold=True, color=COLORS["red"] if "atrasados" in label else COLORS["text"])
        ws.cell(6, col).number_format = fmt
        for row in [5, 6]:
            for c in [col, col + 1]:
                ws.cell(row, c).fill = PatternFill("solid", fgColor=COLORS["white"])
                ws.cell(row, c).border = Border(left=Side(style="thin", color=COLORS["line"]), right=Side(style="thin", color=COLORS["line"]), top=Side(style="thin", color=COLORS["line"]), bottom=Side(style="thin", color=COLORS["line"]))
    write_small_table(ws, "B9", "Projetos por status", summary["graficos"]["status"])
    write_small_table(ws, "E9", "Projetos por tipo", summary["graficos"]["tipo_projeto"], 8)
    write_small_table(ws, "I9", "Projetos por responsável", summary["graficos"]["responsavel"], 8)
    write_small_table(ws, "B25", "Pendências por categoria", summary["graficos"]["categoria_pendencia"], 8)
    write_small_table(ws, "E25", "Valor por tipo", summary["graficos"]["valor_tipo"], 8)
    ws["I24"] = "Leituras executivas"
    style_header(ws["I24"])
    for row, item in enumerate(summary["insights"][:6], start=25):
        ws[f"I{row}"] = item
        ws[f"I{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, 15):
        ws.column_dimensions[chr(64 + col)].width = 15
    ws.column_dimensions["I"].width = 38
    ws.freeze_panes = "B8"


def add_charts(ws) -> None:
    for title, source, target in [
        ("Projetos por status", "B10:C18", "B36"),
        ("Projetos por tipo", "E10:F18", "I36"),
        ("Pendências por categoria", "B26:C34", "B54"),
    ]:
        min_cell, max_cell = source.split(":")
        min_col, min_row = ws[min_cell].column, ws[min_cell].row
        max_col, max_row = ws[max_cell].column, ws[max_cell].row
        chart = PieChart() if "Pendências" in title else BarChart()
        chart.add_data(Reference(ws, min_col=max_col, min_row=min_row + 1, max_row=max_row), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=min_col, min_row=min_row + 2, max_row=max_row))
        chart.title = title
        chart.height = 8
        chart.width = 13
        ws.add_chart(chart, target)


def write_dataframe_sheet(ws, df: pd.DataFrame, name: str) -> None:
    ws.sheet_properties.tabColor = COLORS["mint"]
    for row in [list(df.columns)] + df.where(pd.notna(df), None).values.tolist():
        ws.append(row)
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        style_header(cell)
    if ws.max_row > 1 and ws.max_column > 1:
        table_name = ("tbl_" + name.replace(" ", "_").replace("-", "_"))[:31]
        table = Table(displayName=table_name, ref=f"A1:{ws.cell(ws.max_row, ws.max_column).coordinate}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)
    for col in ws.columns:
        letter = col[0].column_letter
        header = str(col[0].value)
        ws.column_dimensions[letter].width = min(36, max(11, max(len(str(cell.value)) if cell.value is not None else 0 for cell in col[:80]) + 2))
        if "Data" in header:
            for cell in col[1:]:
                cell.number_format = "dd/mm/yyyy"
        if header in ["Valor_Projeto", "Custo_Estimado", "Margem_Estimada", "Valor_Faturado"]:
            for cell in col[1:]:
                cell.number_format = '"R$"#,##0.00'
        if header == "Potencia_kW":
            for cell in col[1:]:
                cell.number_format = "#,##0.0"
    if "Status_Simplificado" in df.columns:
        col = list(df.columns).index("Status_Simplificado") + 1
        rng = f"{ws.cell(2, col).coordinate}:{ws.cell(ws.max_row, col).coordinate}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{ws.cell(2, col).coordinate}="ATRASADO"'], fill=PatternFill("solid", fgColor="F9D7D4")))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{ws.cell(2, col).coordinate}="CONCLUÍDO"'], fill=PatternFill("solid", fgColor="DDEBE4")))
    if "Dias_Atraso" in df.columns:
        col = list(df.columns).index("Dias_Atraso") + 1
        rng = f"{ws.cell(2, col).coordinate}:{ws.cell(ws.max_row, col).coordinate}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FCE8B2")))


def write_excel(tables: dict[str, pd.DataFrame], summary: dict[str, object]) -> None:
    wb = Workbook()
    wb.active.title = "01_Dashboard"
    setup_dashboard(wb["01_Dashboard"], summary)
    sheet_map = [
        ("02_Projetos", tables["Fato_Projetos"]),
        ("03_Pendencias", tables["Fato_Pendencias"]),
        ("04_Historico", tables["Fato_HistoricoStatus"]),
        ("05_Horas", tables["Fato_Horas"]),
        ("06_Faturamento", tables["Fato_Faturamento"]),
        ("07_Clientes", tables["Dim_Cliente"]),
        ("08_Responsaveis", tables["Dim_Responsavel"]),
        ("09_Concessionarias", tables["Dim_Concessionaria"]),
        ("10_Tipos_Projeto", tables["Dim_TipoProjeto"]),
        ("11_Status", tables["Dim_Status"]),
        ("12_Etapas", tables["Dim_Etapa"]),
        ("13_Localidades", tables["Dim_Localidade"]),
        ("14_Calendario", tables["Dim_Calendario"]),
    ]
    for name, df in sheet_map:
        write_dataframe_sheet(wb.create_sheet(name), df, name)
    add_charts(wb["01_Dashboard"])
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.save(ROOT / "dashboard_data.xlsx")


def table_md(headers: list[str], rows: list[list[object]]) -> str:
    return "\n".join(["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"] + ["| " + " | ".join(str(item) for item in row) + " |" for row in rows])


def write_powerbi_docs(summary: dict[str, object]) -> None:
    (POWERBI / "theme.json").write_text(json.dumps({"name": "Gestão Projetos Elétricos", "dataColors": [hx("green"), hx("blue"), hx("gold"), hx("amber"), hx("red"), "#7A8691"], "background": hx("white"), "foreground": hx("text"), "tableAccent": hx("green")}, ensure_ascii=False, indent=2), encoding="utf-8")
    dax = """# Medidas DAX

Use `Parametros[DataReferencia]` como data fixa da demonstração.

```DAX
Data Referencia = SELECTEDVALUE(Parametros[DataReferencia], DATE(2026, 9, 8))
Total Projetos = DISTINCTCOUNT(Fato_Projetos[ID_Projeto])
Projetos Ativos = CALCULATE([Total Projetos], FILTER(Fato_Projetos, NOT (Fato_Projetos[Status_Simplificado] IN {"CONCLUÍDO", "SUSPENSO", "CANCELADO"})))
Projetos Concluidos = CALCULATE([Total Projetos], Fato_Projetos[Status_Simplificado] = "CONCLUÍDO")
Projetos Atrasados = CALCULATE([Total Projetos], Fato_Projetos[Status_Simplificado] = "ATRASADO")
Projetos Suspensos = CALCULATE([Total Projetos], Fato_Projetos[Status_Simplificado] = "SUSPENSO")
Projetos Cancelados = CALCULATE([Total Projetos], Fato_Projetos[Status_Simplificado] = "CANCELADO")
Potencia Total kW = SUM(Fato_Projetos[Potencia_kW])
Potencia Total MW = DIVIDE([Potencia Total kW], 1000)
Valor da Carteira = SUM(Fato_Projetos[Valor_Projeto])
Valor Projetos Ativos = CALCULATE([Valor da Carteira], FILTER(Fato_Projetos, NOT (Fato_Projetos[Status_Simplificado] IN {"CONCLUÍDO", "SUSPENSO", "CANCELADO"})))
Ticket Medio = AVERAGE(Fato_Projetos[Valor_Projeto])
Prazo Medio = AVERAGEX(Fato_Projetos, DATEDIFF(Fato_Projetos[Data_Entrada], Fato_Projetos[Data_Prazo], DAY))
% Projetos Concluidos = DIVIDE([Projetos Concluidos], [Total Projetos])
% Concluido no Prazo = DIVIDE(CALCULATE([Total Projetos], FILTER(Fato_Projetos, Fato_Projetos[Status_Simplificado] = "CONCLUÍDO" && Fato_Projetos[Data_Conclusao] <= Fato_Projetos[Data_Prazo])), [Projetos Concluidos])
% Projetos Atrasados = DIVIDE([Projetos Atrasados], [Total Projetos])
Total Pendencias = COUNTROWS(Fato_Pendencias)
Projetos com Pendencia = CALCULATE([Total Projetos], Fato_Projetos[Possui_Pendencia] = TRUE())
Tempo Medio Parado = AVERAGE(Fato_Projetos[Dias_Parado])
Horas Estimadas = SUM(Fato_Projetos[Horas_Estimadas])
Horas Realizadas = SUM(Fato_Projetos[Horas_Realizadas])
Desvio de Horas = [Horas Realizadas] - [Horas Estimadas]
Eficiencia de Horas = DIVIDE([Horas Estimadas], [Horas Realizadas])
Margem Estimada = SUM(Fato_Projetos[Margem_Estimada])
% Margem Estimada = DIVIDE([Margem Estimada], [Valor da Carteira])
```
"""
    (POWERBI / "dax_measures.md").write_text(dax, encoding="utf-8")
    (POWERBI / "power_query.md").write_text("""# Power Query

Importe todos os CSVs da pasta `data` com UTF-8, promova cabeçalhos e tipifique datas, moedas, potência, horas e campos booleanos.

```powerquery
let
    Fonte = Csv.Document(File.Contents("data/Fato_Projetos.csv"), [Delimiter=",", Encoding=65001, QuoteStyle=QuoteStyle.Csv]),
    Cabecalhos = Table.PromoteHeaders(Fonte, [PromoteAllScalars=true]),
    Tipos = Table.TransformColumnTypes(Cabecalhos, {{"ID_Projeto", type text}, {"Data_Entrada", type date}, {"Data_Prazo", type date}, {"Potencia_kW", type number}, {"Valor_Projeto", Currency.Type}})
in
    Tipos
```
""", encoding="utf-8")
    relationships = [["Dim_Cliente[ID_Cliente]", "Fato_Projetos[ID_Cliente]"], ["Dim_Responsavel[ID_Responsavel]", "Fato_Projetos[ID_Responsavel]"], ["Dim_Concessionaria[ID_Concessionaria]", "Fato_Projetos[ID_Concessionaria]"], ["Dim_TipoProjeto[ID_TipoProjeto]", "Fato_Projetos[ID_TipoProjeto]"], ["Dim_Status[ID_Status]", "Fato_Projetos[ID_Status]"], ["Dim_Etapa[ID_Etapa]", "Fato_Projetos[ID_Etapa]"], ["Dim_Localidade[ID_Localidade]", "Fato_Projetos[ID_Localidade]"], ["Dim_Calendario[Data]", "Fato_Projetos[Data_Entrada]"], ["Fato_Projetos[ID_Projeto]", "Fato_Pendencias[ID_Projeto]"], ["Fato_Projetos[ID_Projeto]", "Fato_Horas[ID_Projeto]"], ["Fato_Projetos[ID_Projeto]", "Fato_Faturamento[ID_Projeto]"]]
    (POWERBI / "data_model.md").write_text("# Modelo de Dados\n\n" + table_md(["Origem", "Destino", "Cardinalidade", "Filtro"], [[a, b, "1:*", "Único"] for a, b in relationships]), encoding="utf-8")
    pages = ["Visão Executiva", "Gestão de Prazos", "Operações", "Engenharia", "Financeiro", "Pendências e Gargalos", "Detalhamento do Projeto"]
    (POWERBI / "dashboard_specification.md").write_text("# Especificação do Dashboard\n\n" + "\n\n".join([f"## Página {idx}: {page}\nFiltros: ano, estado, concessionária, tipo de projeto, responsável, status, etapa e prioridade. Visualizações: cartões, barras, tabela analítica e leitura executiva." for idx, page in enumerate(pages, 1)]), encoding="utf-8")
    (POWERBI / "GUIA_MONTAGEM_POWER_BI.md").write_text(f"""# Guia de Montagem no Power BI

Não foi gerado PBIX automático porque Power BI Desktop, Tabular Editor CLI e pbi-tools não estavam disponíveis no ambiente de finalização.

1. Abra o Power BI Desktop.
2. Importe os CSVs da pasta `data` com UTF-8.
3. Monte os relacionamentos de `powerbi/data_model.md`.
4. Cole as medidas de `powerbi/dax_measures.md`.
5. Aplique o tema `powerbi/theme.json`.
6. Crie as sete páginas descritas em `powerbi/dashboard_specification.md`.
7. Use `assets/mockups` como referência visual.
8. Confira os cartões contra `outputs/resumo_portfolio.json`.

Conferência: {summary['metricas_formatadas']['total_projetos']} projetos, {summary['metricas_formatadas']['projetos_atrasados']} atrasados, {summary['metricas_formatadas']['projetos_concluidos']} concluídos, {summary['metricas_formatadas']['total_pendencias']} pendências, {summary['metricas_formatadas']['potencia_total_mw']} e {summary['metricas_formatadas']['valor_carteira']}.
""", encoding="utf-8")
    (POWERBI / "pbip_notes.md").write_text("# Observações sobre PBIX/PBIP\n\nA entrega não contém PBIX simulado. Use o guia de montagem para criar o arquivo oficial no Power BI Desktop.\n", encoding="utf-8")


def field_desc(col: str) -> str:
    mapping = {"ID_Projeto": "Identificador único do projeto.", "Nome_Projeto": "Nome fictício do projeto.", "Potencia_kW": "Potência elétrica associada ao projeto.", "Valor_Projeto": "Valor comercial estimado.", "Status_Simplificado": "Status executivo calculado.", "Dias_Atraso": "Dias de atraso na data de referência.", "Municipio": "Município de referência sintético."}
    return mapping.get(col, col.replace("_", " ").capitalize())


def write_docs(summary: dict[str, object], tables: dict[str, pd.DataFrame]) -> None:
    kf = summary["metricas_formatadas"]
    field_rows = []
    for name in sorted(tables):
        sample = tables[name].iloc[0].to_dict() if len(tables[name]) else {}
        for col in tables[name].columns:
            field_rows.append([name, col, str(tables[name][col].dtype), field_desc(col), sample.get(col, "")])
    (DOCS / "data_dictionary.md").write_text("# Dicionário de Dados\n\n" + table_md(["Tabela", "Campo", "Tipo", "Descrição", "Exemplo"], field_rows), encoding="utf-8")
    kpis = [["Total de Projetos", "Contagem distinta de projetos.", kf["total_projetos"]], ["Projetos Ativos", "Projetos não encerrados.", kf["projetos_ativos"]], ["Projetos Atrasados", "Projetos abertos com prazo vencido.", kf["projetos_atrasados"]], ["Potência Total", "Soma da potência em MW.", kf["potencia_total_mw"]], ["Valor da Carteira", "Soma dos valores de projeto.", kf["valor_carteira"]], ["Conclusão no Prazo", "Concluídos até a data de prazo.", kf["taxa_concluido_prazo"]], ["Total de Pendências", "Registros de pendência.", kf["total_pendencias"]], ["Eficiência de Horas", "Horas estimadas divididas por realizadas.", kf["eficiencia_horas"]]]
    (DOCS / "kpi_dictionary.md").write_text("# Dicionário de KPIs\n\n" + table_md(["KPI", "Definição", "Valor Atual"], kpis), encoding="utf-8")
    (DOCS / "insights.md").write_text("# Insights Calculados\n\n" + "\n".join(f"- {item}" for item in summary["insights"]) + "\n", encoding="utf-8")
    (DOCS / "executive_summary.md").write_text(f"# Resumo Executivo\n\nA carteira sintética contém {kf['total_projetos']} projetos, {kf['potencia_total_mw']} e {kf['valor_carteira']}. Há {kf['projetos_ativos']} ativos, {kf['projetos_atrasados']} atrasados e {kf['total_pendencias']} pendências.\n", encoding="utf-8")
    (DOCS / "project_overview.md").write_text("# Visão Geral do Projeto\n\nProjeto de portfólio que simula a gestão de projetos elétricos com dados sintéticos, modelo estrela, DAX, Excel, mockups Power BI e demo web.\n", encoding="utf-8")
    portfolio = "# Descrição para Portfólio\n\nDashboard de Gestão de Projetos Elétricos criado para demonstrar competências em engenharia elétrica, Power BI, Python, modelagem de dados, DAX, Power Query, análise de prazos, gestão de pendências e comunicação executiva.\n"
    (DOCS / "descricao_portfolio.md").write_text(portfolio, encoding="utf-8")
    (DOCS / "portfolio_description.md").write_text(portfolio, encoding="utf-8")
    questions = ["Qual problema de negócio o dashboard resolve?", "Por que usar modelo estrela?", "Como o status ATRASADO foi calculado?", "Como validar os prazos?", "Por que separar fatos e dimensões?", "Quais filtros são mais importantes?", "Como priorizar projetos críticos?", "Como tratar dados reais de concessionárias?", "Como evitar dupla contagem de pendências?", "Qual KPI mais ajuda a engenharia?", "Qual KPI mais ajuda a diretoria?", "Como explicar MW para público executivo?", "Como melhorar a previsão de horas?", "Quais riscos existem em dados sintéticos?", "Como versionar no GitHub?", "Como publicar no GitHub Pages?", "Como garantir privacidade?", "Como expandir para múltiplas empresas?", "Como integrar ERP ou CRM?", "O que faria em uma segunda versão?"]
    (DOCS / "interview_case.md").write_text("# Case para Entrevistas\n\n" + "\n".join(f"{i}. {q}" for i, q in enumerate(questions, 1)) + "\n", encoding="utf-8")
    (DOCS / "curriculo_projeto.md").write_text(f"""# Currículo e ATS

## Versão curta
Criei um dashboard de gestão de projetos elétricos com Python, Power BI, DAX, Power Query, Excel e demo web.

## Versão média
Desenvolvi um projeto completo de BI para gestão de projetos de engenharia elétrica, cobrindo prazos, pendências, potência, concessionárias, responsáveis, custos, margem e status executivo.

## Versão ATS
Power BI, DAX, Power Query, Python, pandas, openpyxl, Excel, modelagem dimensional, modelo estrela, engenharia elétrica, gestão de projetos, análise de prazos, análise financeira, dashboards executivos, GitHub Pages, documentação técnica.

## Resultados
- {kf['total_projetos']} projetos sintéticos validados.
- {kf['potencia_total_mw']} de potência analisada.
- {kf['valor_carteira']} em valor de carteira.
- {kf['total_pendencias']} pendências operacionais mapeadas.
""", encoding="utf-8")


def write_linkedin(summary: dict[str, object], repo_url: str = REPO_URL, dashboard_url: str = DASHBOARD_URL) -> None:
    k = summary["metricas_formatadas"]
    files = {
        "01_PROJETO_LINKEDIN.txt": f"Dashboard de Gestão de Projetos Elétricos\n\nProjeto completo em Power BI e Python para gestão de uma carteira sintética de projetos elétricos.\n\nIndicadores: {k['total_projetos']} projetos, {k['potencia_total_mw']}, {k['valor_carteira']}, {k['projetos_atrasados']} atrasados e {k['total_pendencias']} pendências.\n",
        "02_COMPETENCIAS.txt": "Power BI\nDAX\nPower Query\nPython\npandas\nopenpyxl\nExcel\nModelagem dimensional\nModelo estrela\nEngenharia elétrica\nGestão de projetos\nAnálise de prazos\nAnálise financeira\nDashboards executivos\nGitHub Pages\nDocumentação técnica\n",
        "03_POST_LINKEDIN.md": f"# Post LinkedIn\n\nFinalizei um projeto completo de BI para gestão de projetos elétricos.\n\nCriei dados sintéticos com Python, modelei uma carteira com {k['total_projetos']} projetos, documentei medidas DAX, preparei Excel profissional, especifiquei sete páginas para Power BI e publiquei uma demo web com filtros.\n\nPrincipais números:\n- Potência analisada: {k['potencia_total_mw']}\n- Valor da carteira: {k['valor_carteira']}\n- Projetos atrasados: {k['projetos_atrasados']}\n- Pendências registradas: {k['total_pendencias']}\n\nRepositório: {repo_url}\nDashboard: {dashboard_url}\n\n#PowerBI #Python #DAX #EngenhariaEletrica #GestaoDeProjetos #BusinessIntelligence #DataAnalytics\n",
        "04_COMO_ADICIONAR.txt": "1. No LinkedIn, abra Adicionar seção do perfil.\n2. Escolha Projetos.\n3. Use o título Dashboard de Gestão de Projetos Elétricos.\n4. Informe a URL do GitHub Pages ou do repositório.\n5. Use o texto de 01_PROJETO_LINKEDIN.txt.\n6. Publique as imagens de capa e arte quadrada.\n",
        "05_URLS.txt": f"Repositório: {repo_url}\nDashboard: {dashboard_url}\nExcel: dashboard_data.xlsx\nResumo JSON: outputs/resumo_portfolio.json\n",
        "06_RESUMO_30_SEGUNDOS.txt": f"Projeto de BI para gestão de projetos elétricos. Usei Python para gerar e validar uma carteira sintética com {k['total_projetos']} projetos, {k['potencia_total_mw']} e {k['valor_carteira']}. A entrega inclui Excel, DAX, Power Query, mockups Power BI, documentação e demo web com filtros.\n",
        "07_RESUMO_1_MINUTO.txt": f"Este projeto simula uma operação real de engenharia elétrica. A base tem {k['total_projetos']} projetos, concessionárias, responsáveis, prazos, etapas, potência, valores e pendências. Eu construí um pipeline reproduzível com Python, gerei um Excel estruturado, documentei o modelo estrela, escrevi medidas DAX, criei mockups para sete páginas de Power BI e implementei uma demo web estática com filtros.\n",
    }
    for filename, content in files.items():
        (LINKEDIN / filename).write_text(content, encoding="utf-8")
    (DOCS / "linkedin_post.md").write_text(files["03_POST_LINKEDIN.md"], encoding="utf-8")
    (DOCS / "linkedin_project.md").write_text(files["01_PROJETO_LINKEDIN.txt"], encoding="utf-8")


def write_readme(summary: dict[str, object], dashboard_url: str = DASHBOARD_URL) -> None:
    k = summary["metricas_formatadas"]
    dashboard_line = f"\nDemo online: [{dashboard_url}]({dashboard_url})\n" if dashboard_url else "\nDemo local: abra `web/index.html` ou execute `INICIAR_DASHBOARD.bat`.\n"
    readme = f"""# Dashboard de Gestão de Projetos Elétricos

Projeto completo de portfólio em Business Intelligence para gestão de projetos de engenharia elétrica, com dados sintéticos, Excel profissional, documentação Power BI, medidas DAX, mockups e demo web estática.
{dashboard_line}
## Prévia do Dashboard

![Visão Executiva](assets/mockups/executive_overview.png)
![Gestão de Prazos](assets/mockups/deadlines.png)
![Operações](assets/mockups/operations.png)
![Engenharia](assets/mockups/engineering.png)
![Financeiro](assets/mockups/financial.png)
![Pendências e Gargalos](assets/mockups/issues.png)
![Detalhamento do Projeto](assets/mockups/project_detail.png)

## Indicadores Principais

| Indicador | Valor |
| --- | --- |
| Total de projetos | {k['total_projetos']} |
| Projetos ativos | {k['projetos_ativos']} |
| Projetos atrasados | {k['projetos_atrasados']} |
| Projetos concluídos | {k['projetos_concluidos']} |
| Potência total | {k['potencia_total_mw']} |
| Valor da carteira | {k['valor_carteira']} |
| Pendências | {k['total_pendencias']} |
| Conclusão no prazo | {k['taxa_concluido_prazo']} |

## Entregas

- `data/`: CSVs em UTF-8 com dados sintéticos.
- `dashboard_data.xlsx`: planilha com dashboard, tabelas, filtros, congelamento de painéis, formatação e gráficos.
- `powerbi/`: DAX, Power Query, modelo, tema e guia de montagem.
- `web/`: dashboard estático com sete páginas e filtros funcionais.
- `docs/`: dicionários, resumo executivo, insights, currículo e case de entrevista.
- `linkedin/`: textos e imagens para publicação no LinkedIn.
- `outputs/resumo_portfolio.json`: fonte única dos KPIs finais.

## Como Atualizar

```bash
python scripts/finalizar_projeto.py
```

No Windows, execute `ATUALIZAR_TUDO.bat`. Para abrir a demo local, execute `INICIAR_DASHBOARD.bat`.

## Power BI

O arquivo PBIX não é simulado. Para montar o relatório oficial, siga `powerbi/GUIA_MONTAGEM_POWER_BI.md`.

## Privacidade

Todos os dados são sintéticos e fictícios. Não há clientes, protocolos, documentos ou informações confidenciais reais.

## Autor

Lucas Almeida
"""
    (ROOT / "README.md").write_text(readme, encoding="utf-8")
    (ROOT / "README_PT-BR.md").write_text("# README em Português\n\nO README principal já está em português do Brasil: [README.md](README.md).\n", encoding="utf-8")


def write_web() -> None:
    (WEB / "index.html").write_text("""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8" /><meta name="viewport" content="width=device-width, initial-scale=1" /><title>Dashboard de Gestão de Projetos Elétricos</title><link rel="stylesheet" href="styles.css" /></head><body><aside class="sidebar"><div class="brand"><strong>Gestão Elétrica</strong><span>Portfólio BI</span></div><nav><button class="active" data-page="executiva">Visão Executiva</button><button data-page="prazos">Gestão de Prazos</button><button data-page="operacoes">Operações</button><button data-page="engenharia">Engenharia</button><button data-page="financeiro">Financeiro</button><button data-page="pendencias">Pendências e Gargalos</button><button data-page="detalhe">Detalhamento do Projeto</button></nav></aside><main><header class="topbar"><div><h1>Dashboard de Gestão de Projetos Elétricos</h1><p id="contexto"></p></div><button id="limpar">Limpar filtros</button></header><section class="filters" id="filters"></section><section class="cards" id="cards"></section><section id="page" class="page"></section></main><script src="portfolio_summary.js"></script><script src="app.js"></script></body></html>""", encoding="utf-8")
    (WEB / "styles.css").write_text("""*{box-sizing:border-box}body{margin:0;font-family:Segoe UI,Arial,sans-serif;background:#f6f8f7;color:#20252b;display:flex;min-height:100vh}.sidebar{width:256px;background:#12372a;color:#fff;padding:26px 18px;position:sticky;top:0;height:100vh}.brand{margin-bottom:24px}.brand strong{display:block;font-size:21px}.brand span{color:#cfe1d7;font-size:13px}nav button,#limpar{font:inherit;border-radius:6px;border:1px solid #315c4b;cursor:pointer}nav button{width:100%;text-align:left;margin:5px 0;padding:11px 12px;background:transparent;color:#fff}nav button.active,nav button:hover{background:#1f5b43}main{flex:1;padding:24px 30px;min-width:0}.topbar{display:flex;align-items:end;justify-content:space-between;gap:20px;margin-bottom:18px}h1{font-size:28px;margin:0 0 4px}h2{font-size:18px;margin:0 0 14px}p{margin:0;color:#69737d}.filters{display:grid;grid-template-columns:repeat(4,minmax(160px,1fr));gap:12px;margin-bottom:18px}.filter label{display:block;font-size:12px;color:#69737d;margin-bottom:4px}.filter select{width:100%;padding:10px;border:1px solid #cfd8d3;border-radius:6px;background:white;color:#20252b}.cards{display:grid;grid-template-columns:repeat(6,minmax(130px,1fr));gap:12px;margin-bottom:18px}.card,.panel{background:#fff;border:1px solid #dde4e0;border-radius:8px}.card{padding:14px;border-left:5px solid #1f5b43;min-height:84px}.card.risk{border-left-color:#b42318}.card span{display:block;color:#69737d;font-size:12px}.card strong{display:block;font-size:22px;margin-top:7px}.page{display:grid;grid-template-columns:1fr 1fr;gap:16px}.panel{padding:17px;min-width:0}.wide{grid-column:1/-1}.bar{display:grid;grid-template-columns:minmax(120px,210px) 1fr 58px;gap:10px;align-items:center;margin:9px 0}.bar span{white-space:nowrap;overflow:hidden;text-overflow:ellipsis;color:#364049}.track{height:15px;background:#edf2ef;border-radius:5px}.fill{height:15px;background:#1f5b43;border-radius:5px}.insights li{margin:8px 0}.table-wrap{overflow:auto}table{width:100%;border-collapse:collapse;min-width:740px}th,td{padding:10px;border-bottom:1px solid #dde4e0;text-align:left;font-size:13px;vertical-align:top}th{background:#12372a;color:#fff;position:sticky;top:0}.badge{display:inline-block;padding:4px 8px;border-radius:999px;background:#ddebe4;color:#12372a;font-weight:600}.badge.risk{background:#f9d7d4;color:#8a1c13}#limpar{padding:10px 14px;background:#fff;color:#12372a;border-color:#cfd8d3}.detail-grid{display:grid;grid-template-columns:repeat(3,1fr);gap:12px}.detail-item{padding:12px;background:#f6f8f7;border:1px solid #dde4e0;border-radius:8px}.detail-item span{display:block;font-size:12px;color:#69737d}.detail-item strong{display:block;margin-top:5px}@media(max-width:1100px){.cards{grid-template-columns:repeat(3,1fr)}.filters{grid-template-columns:repeat(2,1fr)}.page{grid-template-columns:1fr}.sidebar{position:relative;height:auto;width:100%}body{display:block}}@media(max-width:650px){main{padding:18px}.topbar{display:block}.cards,.filters,.detail-grid{grid-template-columns:1fr}h1{font-size:23px}.card strong{font-size:20px}}""", encoding="utf-8")
    (WEB / "app.js").write_text("""const data=window.PORTFOLIO_SUMMARY;let state={page:'executiva',filters:{}};const fmtInt=v=>Math.round(Number(v)||0).toLocaleString('pt-BR');const fmtDec=(v,p=1)=>(Number(v)||0).toLocaleString('pt-BR',{minimumFractionDigits:p,maximumFractionDigits:p});const fmtMoney=v=>(Number(v)||0).toLocaleString('pt-BR',{style:'currency',currency:'BRL',maximumFractionDigits:0});const filterDefs=[['ano','Ano','Data_Entrada',p=>(p.Data_Entrada||'').slice(0,4)],['estado','Estado','Estado',p=>p.Estado],['concessionaria','Concessionária','Concessionaria',p=>p.Concessionaria],['tipo','Tipo de Projeto','Tipo_Projeto',p=>p.Tipo_Projeto],['responsavel','Responsável','Responsavel',p=>p.Responsavel],['status','Status','Status_Simplificado',p=>p.Status_Simplificado],['etapa','Etapa','Macro_Etapa',p=>p.Macro_Etapa],['prioridade','Prioridade','Prioridade',p=>p.Prioridade]];function filtered(){return data.projetos.filter(p=>filterDefs.every(([id,,field,getter])=>!state.filters[id]||String(getter(p))===state.filters[id]))}function group(rows,field,measure='count'){const out={};rows.forEach(r=>{const key=r[field]||'n.a.';out[key]=(out[key]||0)+(measure==='count'?1:Number(r[measure]||0))});return Object.fromEntries(Object.entries(out).sort((a,b)=>b[1]-a[1]).slice(0,12))}function metrics(rows){const total=rows.length,concluidos=rows.filter(p=>p.Status_Simplificado==='CONCLUÍDO').length,ativos=rows.filter(p=>!['CONCLUÍDO','CANCELADO','SUSPENSO'].includes(p.Status_Simplificado)).length,atrasados=rows.filter(p=>p.Status_Simplificado==='ATRASADO').length,valor=rows.reduce((s,p)=>s+Number(p.Valor_Projeto||0),0),potencia=rows.reduce((s,p)=>s+Number(p.Potencia_kW||0),0)/1000,pend=rows.reduce((s,p)=>s+Number(p.Quantidade_Pendencias||0),0);return{total,ativos,atrasados,concluidos,valor,potencia,pend}}function renderFilters(){document.querySelector('#filters').innerHTML=filterDefs.map(([id,label,,getter])=>{const opts=[...new Set(data.projetos.map(getter).filter(Boolean).map(String))].sort((a,b)=>a.localeCompare(b,'pt-BR'));return`<div class="filter"><label>${label}</label><select data-filter="${id}"><option value="">Todos</option>${opts.map(o=>`<option ${state.filters[id]===o?'selected':''}>${o}</option>`).join('')}</select></div>`}).join('');document.querySelectorAll('select[data-filter]').forEach(s=>s.onchange=e=>{state.filters[e.target.dataset.filter]=e.target.value;render()})}function renderCards(rows){const m=metrics(rows);const cards=[['Projetos',fmtInt(m.total),''],['Ativos',fmtInt(m.ativos),''],['Atrasados',fmtInt(m.atrasados),'risk'],['Potência',`${fmtDec(m.potencia,1)} MW`,''],['Carteira',fmtMoney(m.valor),''],['Pendências',fmtInt(m.pend),'risk']];document.querySelector('#cards').innerHTML=cards.map(c=>`<div class="card ${c[2]}"><span>${c[0]}</span><strong>${c[1]}</strong></div>`).join('')}function bars(title,obj,format=v=>fmtInt(v)){const max=Math.max(1,...Object.values(obj).map(Number));return`<section class="panel"><h2>${title}</h2>${Object.entries(obj).map(([k,v])=>`<div class="bar" title="${k}"><span>${k}</span><div class="track"><div class="fill" style="width:${Math.max(3,Number(v)/max*100)}%"></div></div><strong>${format(v)}</strong></div>`).join('')}</section>`}function table(title,rows,cols){return`<section class="panel wide"><h2>${title}</h2><div class="table-wrap"><table><thead><tr>${cols.map(c=>`<th>${c[0]}</th>`).join('')}</tr></thead><tbody>${rows.map(r=>`<tr>${cols.map(c=>`<td>${c[2]?c[2](r[c[1]],r):r[c[1]]}</td>`).join('')}</tr>`).join('')}</tbody></table></div></section>`}function detail(project){if(!project)return'<section class="panel wide"><h2>Nenhum projeto encontrado</h2></section>';const fields=[['Projeto',project.Nome_Projeto],['Cliente',project.Cliente],['Município',`${project.Municipio} / ${project.Estado}`],['Concessionária',project.Concessionaria],['Responsável',project.Responsavel],['Status',project.Status_Simplificado],['Tipo',project.Tipo_Projeto],['Etapa',project.Etapa],['Potência',`${fmtDec(project.Potencia_kW,1)} kW`],['Valor',fmtMoney(project.Valor_Projeto)],['Prazo',project.Data_Prazo],['Pendências',fmtInt(project.Quantidade_Pendencias)]];return`<section class="panel wide"><h2>${project.Nome_Projeto}</h2><div class="detail-grid">${fields.map(f=>`<div class="detail-item"><span>${f[0]}</span><strong>${f[1]}</strong></div>`).join('')}</div></section>`}function renderPage(rows){const critical=[...rows].sort((a,b)=>(b.Dias_Atraso-a.Dias_Atraso)||(b.Quantidade_Pendencias-a.Quantidade_Pendencias)).slice(0,12);const pages={executiva:()=>bars('Projetos por status',group(rows,'Status_Simplificado'))+bars('Projetos por tipo',group(rows,'Tipo_Projeto'))+table('Projetos críticos',critical,[['Projeto','Nome_Projeto'],['Responsável','Responsavel'],['Concessionária','Concessionaria'],['Status','Status_Simplificado',v=>`<span class="badge ${v==='ATRASADO'?'risk':''}">${v}</span>`],['Atraso','Dias_Atraso']])+`<section class="panel wide"><h2>Insights calculados</h2><ul class="insights">${data.insights.map(i=>`<li>${i}</li>`).join('')}</ul></section>`,prazos:()=>bars('Faixas de prazo',group(rows,'Faixa_Prazo'))+bars('Atrasos por concessionária',group(rows.filter(p=>p.Status_Simplificado==='ATRASADO'),'Concessionaria'))+table('Mapa de prazos',critical,[['Projeto','Nome_Projeto'],['Prazo','Data_Prazo'],['Faixa','Faixa_Prazo'],['Dias de atraso','Dias_Atraso'],['Criticidade','Criticidade']]),operacoes:()=>bars('Projetos por etapa',group(rows,'Macro_Etapa'))+bars('Carga por responsável',group(rows,'Responsavel'))+table('Operação por responsável',critical,[['Projeto','Nome_Projeto'],['Responsável','Responsavel'],['Etapa','Etapa'],['Horas','Horas_Realizadas',(v,r)=>`${fmtInt(r.Horas_Realizadas)} / ${fmtInt(r.Horas_Estimadas)}`],['Dias parado','Dias_Parado']]),engenharia:()=>bars('Potência por categoria (kW)',group(rows,'Categoria_Engenharia','Potencia_kW'),v=>fmtDec(v,1))+bars('Projetos por estado',group(rows,'Estado'))+table('Carteira técnica',critical,[['Projeto','Nome_Projeto'],['Tipo','Tipo_Projeto'],['Nível','Nivel_Tensao'],['Conexão','Tipo_Conexao'],['Potência kW','Potencia_kW',v=>fmtDec(v,1)]]),financeiro:()=>bars('Valor por tipo',group(rows,'Tipo_Projeto','Valor_Projeto'),fmtMoney)+bars('Margem por categoria',group(rows,'Categoria_Engenharia','Margem_Estimada'),fmtMoney)+table('Projetos por valor',[...rows].sort((a,b)=>b.Valor_Projeto-a.Valor_Projeto).slice(0,12),[['Projeto','Nome_Projeto'],['Cliente','Cliente'],['Valor','Valor_Projeto',fmtMoney],['Custo','Custo_Estimado',fmtMoney],['Margem','Margem_Estimada',fmtMoney]]),pendencias:()=>bars('Pendências por categoria',group(data.pendencias.filter(p=>rows.some(r=>r.ID_Projeto===p.ID_Projeto)),'Categoria_Pendencia'))+bars('Projetos por criticidade',group(rows,'Criticidade'))+table('Gargalos',critical,[['Projeto','Nome_Projeto'],['Pendências','Quantidade_Pendencias'],['Dias parado','Dias_Parado'],['Observação','Observacao_Status'],['Responsável','Responsavel']]),detalhe:()=>detail(rows[0])};document.querySelector('#page').innerHTML=(pages[state.page]||pages.executiva)()}function render(){const rows=filtered();document.querySelector('#contexto').textContent=`${fmtInt(rows.length)} projetos filtrados | referência ${data.metadata.data_referencia}`;renderFilters();renderCards(rows);renderPage(rows)}document.querySelectorAll('nav button').forEach(btn=>btn.onclick=()=>{document.querySelectorAll('nav button').forEach(b=>b.classList.remove('active'));btn.classList.add('active');state.page=btn.dataset.page;render()});document.querySelector('#limpar').onclick=()=>{state.filters={};render()};render();""", encoding="utf-8")


def write_scripts() -> None:
    (ROOT / "ATUALIZAR_TUDO.bat").write_text('@echo off\nsetlocal\ncd /d "%~dp0"\npython scripts\\finalizar_projeto.py\npause\n', encoding="utf-8")
    (ROOT / "INICIAR_DASHBOARD.bat").write_text('@echo off\nsetlocal\ncd /d "%~dp0web"\necho Abrindo dashboard em http://localhost:8000\nstart "Dashboard Projetos Eletricos" /min python -m http.server 8000\ntimeout /t 2 /nobreak >nul\nstart "" "http://localhost:8000"\n', encoding="utf-8")
    (ROOT / "PUBLICAR_NO_GITHUB.bat").write_text('@echo off\nsetlocal\ncd /d "%~dp0"\ncall ATUALIZAR_TUDO.bat\ngit init\ngit branch -M main\ngit add .\ngit commit -m "Finaliza dashboard de gestao de projetos eletricos" || echo Nada novo para commit.\ngh repo view renovera1/electrical-project-management-dashboard >nul 2>&1 || gh repo create renovera1/electrical-project-management-dashboard --public --source=. --remote=origin\ngit remote get-url origin >nul 2>&1 || git remote add origin https://github.com/renovera1/electrical-project-management-dashboard.git\ngit push -u origin main\ngh pages deploy web --repo renovera1/electrical-project-management-dashboard --branch gh-pages\necho.\necho Repositorio: https://github.com/renovera1/electrical-project-management-dashboard\necho Dashboard: https://renovera1.github.io/electrical-project-management-dashboard/\npause\n', encoding="utf-8")


def write_reports(summary: dict[str, object], github_status: str = "Aguardando publicação") -> None:
    k = summary["metricas_formatadas"]
    inventory = [str(p.relative_to(ROOT)).replace("\\", "/") for p in sorted(ROOT.rglob("*")) if p.is_file() and ".git" not in p.parts and "__pycache__" not in p.parts]
    (ROOT / "INVENTARIO_FINAL.txt").write_text("\n".join(inventory) + "\n", encoding="utf-8")
    (ROOT / "RELATORIO_FINAL.md").write_text(f"""# Relatório Final

## Status
PROJETO FINALIZADO para portfólio, LinkedIn, Excel, Power BI e GitHub Pages.

## KPIs Validados
- Total de projetos: {k['total_projetos']}
- Projetos ativos: {k['projetos_ativos']}
- Projetos atrasados: {k['projetos_atrasados']}
- Projetos concluídos: {k['projetos_concluidos']}
- Potência total: {k['potencia_total_mw']}
- Valor da carteira: {k['valor_carteira']}
- Pendências: {k['total_pendencias']}

## Entregas
- README principal em português.
- Demo web estática com sete páginas e filtros funcionais.
- Excel com dashboard, tabelas, filtros, congelamento de painéis, formatação e gráficos.
- Documentação Power BI com DAX, Power Query, modelo, tema e guia de montagem.
- Materiais de LinkedIn, currículo, portfólio e entrevista.
- ZIP final e hash SHA-256.

## GitHub
{github_status}
""", encoding="utf-8")


def zip_final() -> tuple[Path, str]:
    zip_path = OUTPUTS / "electrical-project-management-dashboard-FINAL-PTBR.zip"
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in ROOT.rglob("*"):
            if path.is_file() and ".git" not in path.parts and "__pycache__" not in path.parts:
                archive.write(path, Path(ROOT.name) / path.relative_to(ROOT))
    sha = hashlib.sha256(zip_path.read_bytes()).hexdigest()
    (OUTPUTS / "electrical-project-management-dashboard-FINAL-PTBR.sha256.txt").write_text(f"{sha}  {zip_path.name}\n", encoding="utf-8")
    return zip_path, sha


def write_all(summary: dict[str, object], tables: dict[str, pd.DataFrame]) -> None:
    write_json(summary)
    write_excel(tables, summary)
    make_mockups(summary)
    make_architecture()
    write_web()
    write_powerbi_docs(summary)
    write_docs(summary, tables)
    write_linkedin(summary)
    write_readme(summary)
    write_scripts()
    write_reports(summary, f"Repositório: {REPO_URL}\nDashboard: {DASHBOARD_URL}")


def clean_temp() -> None:
    for path in ROOT.rglob("__pycache__"):
        shutil.rmtree(path, ignore_errors=True)
    for path in ROOT.rglob("*.pyc"):
        path.unlink(missing_ok=True)


def main() -> None:
    ensure_dirs()
    tables = read_tables()
    df = join_projects(tables)
    summary = build_summary(df, tables["Fato_Pendencias"], tables)
    write_all(summary, tables)
    clean_temp()
    zip_path, sha = zip_final()
    print("Ativos finais gerados em PT-BR.")
    print(f"Projetos: {summary['metricas']['total_projetos']} | Pendências: {summary['metricas']['total_pendencias']} | Atrasados: {summary['metricas']['projetos_atrasados']}")
    print(f"ZIP final: {zip_path}")
    print(f"SHA-256: {sha}")


if __name__ == "__main__":
    main()
